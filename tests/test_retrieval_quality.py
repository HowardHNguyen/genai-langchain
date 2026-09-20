"""Synthetic regression fixtures only: never commit user documents."""
import io
import unittest
from html.parser import HTMLParser
from zipfile import ZipFile, ZIP_STORED
from docx import Document as WordDocument
from langchain_core.documents import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from answer_rendering import answer_html
from document_loader import load_document
from limits import MAX_FILE_BYTES, MAX_PDF_PAGES
from rag import make_prompt, token_count, MAX_INPUT_TOKENS, ask, NO_EVIDENCE, normalize_citations
from retriever import DocumentRetriever, Upload, KeywordIndex
from test_app import FakeEmbeddings, FakeModel


def overview_docx():
    doc = WordDocument()
    doc.add_paragraph('THE 10 OBSERVATORY RESEARCH DOMAINS')
    doc.add_paragraph('The institute coordinates ten domains of research.')
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ['Domain', 'Name', 'Purpose']):
        cell.text = text
    for i in range(1, 11):
        for cell, text in zip(table.add_row().cells, [str(i), f'Area {i}', f'Unique goal {i}. ' * 10]):
            cell.text = text
    for i in range(25):
        doc.add_heading(f'Section {i}: Detailed experiments', level=2)
        doc.add_paragraph('An individual research experiment is described here. ' * 50)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class RetrievalQualityTests(unittest.TestCase):
    def test_overview_retains_all_ten_rows_despite_poor_semantic_scores(self):
        r = DocumentRetriever(FakeEmbeddings())
        r.build([Upload('synthetic.docx', overview_docx())])
        for query in ('what are the 10 observatory research domains?', 'THE TEN OBSERVATORY RESEARCH DOMAINS'):
            docs = r.invoke(query)
            self.assertEqual(docs[0].metadata['block_type'], 'table')
            self.assertIn('THE 10 OBSERVATORY RESEARCH DOMAINS', docs[0].page_content)
            messages, sources, _ = make_prompt(query, [], docs)
            for i in range(1, 11):
                self.assertIn(f'Area {i} | Unique goal {i}.', sources[0]['text'])
            self.assertLessEqual(sum(token_count(m.content) + 8 for m in messages), MAX_INPUT_TOKENS)

    def test_keyword_search_handles_punctuation_and_plural(self):
        index = KeywordIndex(['other item', 'The ten research domains'])
        self.assertEqual(index.search('10 research domain?')[0], 1)

    def test_batching_progress_and_rollback_after_late_failure(self):
        class FailingAfterFirst(FakeEmbeddings):
            def embed_documents(self, texts):
                if self.calls:
                    raise RuntimeError('failed late')
                return super().embed_documents(texts)
        r = DocumentRetriever(FakeEmbeddings())
        r.build([Upload('old.txt', b'apple original')])
        original = r.store
        r.embeddings = FailingAfterFirst()
        updates = []
        with self.assertRaises(RuntimeError):
            r.build([Upload('synthetic.docx', overview_docx())], progress=lambda v, m: updates.append(v))
        self.assertIs(r.store, original)
        self.assertTrue(any(v > .15 for v in updates))

    def test_complete_table_fits_bounded_prompt_with_long_history(self):
        doc = Document(page_content='row value ' * 800, metadata={'block_type': 'table'})
        messages, sources, selected = make_prompt('List rows', [('q', 'a ' * 500)] * 6, [doc] * 20)
        self.assertLessEqual(sum(token_count(m.content) + 8 for m in messages), MAX_INPUT_TOKENS)
        self.assertTrue(sources)
        self.assertEqual(sources[0]['text'], doc.page_content)
        self.assertEqual(len(sources), len(selected))

    def test_one_broader_retry_only_on_new_evidence(self):
        first = Document(page_content='Overview fragment', metadata={'parent_id': 1})
        second = Document(page_content='Complete overview', metadata={'parent_id': 2})
        class Search:
            def invoke(self, query, expanded=False):
                return [first, second] if expanded else [first]
        model = FakeModel([NO_EVIDENCE, 'Complete answer [2]'])
        result = ask(Search(), model, 'List all areas', [])
        self.assertEqual(len(model.calls), 2)
        self.assertEqual(result['sources'][0]['text'], 'Complete overview')

    def test_no_repeated_generation_when_search_cannot_expand(self):
        r = DocumentRetriever(FakeEmbeddings())
        r.build([Upload('small.txt', b'apple')])
        model = FakeModel([NO_EVIDENCE])
        self.assertEqual(ask(r, model, 'What is missing?', [])['answer'], NO_EVIDENCE)
        self.assertEqual(len(model.calls), 1)


class LargeDocumentTests(unittest.TestCase):
    def test_size_and_extracted_text_limits_raised(self):
        self.assertEqual(MAX_FILE_BYTES, 100 * 1024 * 1024)
        buffer = io.BytesIO(overview_docx())
        with ZipFile(buffer, 'a') as archive:
            archive.writestr('word/media/synthetic-padding.bin', b'0' * (11 * 1024 * 1024), compress_type=ZIP_STORED)
        self.assertGreater(len(buffer.getvalue()), 10 * 1024 * 1024)
        self.assertTrue(load_document('large-media.docx', buffer.getvalue()))
        data = ('apple policy\n' * 50000).encode()
        self.assertGreater(len(data), 300000)
        self.assertTrue(load_document('long.txt', data))

    def test_pdf_over_300_pages_including_late_page_evidence(self):
        writer = PdfWriter()
        for i in range(501):
            page = writer.add_blank_page(width=612, height=792)
            if i == 500:
                font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
                page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
                stream = DecodedStreamObject()
                stream.set_data(b'BT /F1 12 Tf 50 700 Td (Late-page verification fact) Tj ET')
                page[NameObject('/Contents')] = writer._add_object(stream)
        buffer = io.BytesIO()
        writer.write(buffer)
        self.assertGreater(MAX_PDF_PAGES, 501)
        docs = load_document('501-pages.pdf', buffer.getvalue())
        self.assertEqual(docs[-1].metadata['page'], 501)
        self.assertIn('verification fact', docs[-1].page_content)


class CitationTests(unittest.TestCase):
    def test_common_reference_formats(self):
        for marker in ('[1]', '[Source 1]', '[S1]', '【1】', '【1†source】', '(Source 1)', '(id 1)'):
            answer, cited, warning = normalize_citations('Fact ' + marker, 3)
            self.assertEqual(cited, {1})
            self.assertEqual(warning, '')
            self.assertEqual(answer, 'Fact [1]')

    def test_grouped_ranges_and_invalid_references(self):
        answer, cited, warning = normalize_citations('Facts [1–3], [1, 99], (ids 2‑3)', 3)
        self.assertEqual(cited, {1, 2, 3})
        self.assertNotIn('[99]', answer)
        self.assertTrue(warning)

    def test_missing_citations_preserve_answer_with_honest_warning(self):
        answer, cited, warning = normalize_citations('Useful answer with no reference.', 2)
        self.assertEqual(answer, 'Useful answer with no reference.')
        self.assertFalse(cited)
        self.assertIn('not as verified support', warning)


class RenderingTests(unittest.TestCase):
    def test_formatting_and_table_render(self):
        html = answer_html('**Ten areas**\n\n| ID | Name |\n|---|---|\n| 1 | Alpha [1] |\n\n1. First\n2. Second')
        for element in ('<strong>', '<table>', '<th>', '<td>', '<ol>'):
            self.assertIn(element, html)

    def test_no_active_remote_images_links_or_raw_html(self):
        class Elements(HTMLParser):
            def __init__(self):
                super().__init__()
                self.tags, self.attrs = [], []
            def handle_starttag(self, tag, attrs):
                self.tags.append(tag)
                self.attrs.extend(attrs)
        malicious = '''![secret](https://evil.example/secret)
![reference][image]
[image]: https://evil.example/tracker
[click](https://evil.example/secret)
<img src="https://evil.example/raw"><script>alert(1)</script>
<iframe src="https://evil.example/frame"></iframe>'''
        parsed = Elements()
        parsed.feed(answer_html(malicious))
        self.assertFalse(set(parsed.tags) & {'img', 'a', 'script', 'iframe', 'svg', 'link'})
        self.assertFalse(any(key in ('src', 'href', 'onerror', 'onclick') for key, value in parsed.attrs))

if __name__ == '__main__':
    unittest.main()
